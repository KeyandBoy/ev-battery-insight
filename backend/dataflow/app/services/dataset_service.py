import csv
import json
import math
import os
from collections import OrderedDict
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from flask import current_app
from werkzeug.utils import secure_filename

from app.extensions import db
from app.models import DataRecord, Dataset


class DatasetValidationError(ValueError):
    pass


SUPPORTED_EXTENSIONS = {
    ".csv": "csv",
    ".json": "json",
    ".xlsx": "excel",
    ".xls": "excel",
}


def validate_upload_file(file_storage):
    if not file_storage or not file_storage.filename:
        raise DatasetValidationError("Please choose a dataset file to upload.")

    ext = Path(file_storage.filename).suffix.lower()
    source_type = SUPPORTED_EXTENSIONS.get(ext)
    if source_type is None:
        raise DatasetValidationError("Only CSV, Excel, and JSON files are supported.")

    return source_type


def save_upload_file(file_storage, user_id):
    upload_root = Path(current_app.config.get("UPLOAD_ROOT", "uploads"))
    if not upload_root.is_absolute():
        upload_root = Path(current_app.root_path).parent / upload_root
    target_dir = upload_root / str(user_id)
    target_dir.mkdir(parents=True, exist_ok=True)

    raw_name = secure_filename(file_storage.filename) or "dataset"
    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    final_name = "{0}_{1}".format(timestamp, raw_name)
    final_path = target_dir / final_name

    file_storage.save(final_path)
    return str(final_path.resolve())


def parse_rows(file_path, source_type):
    if source_type == "csv":
        return _parse_csv(file_path)
    if source_type == "json":
        return _parse_json(file_path)
    if source_type == "excel":
        return _parse_excel(file_path)
    raise DatasetValidationError("Unknown dataset source type.")


def _parse_csv(file_path):
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            with open(file_path, "r", encoding=encoding, newline="") as csv_file:
                reader = csv.DictReader(csv_file)
                return [_normalize_raw_row(item) for item in reader if item]
        except UnicodeDecodeError:
            continue
    raise DatasetValidationError("CSV encoding is not supported. Please use UTF-8.")


def _parse_json(file_path):
    with open(file_path, "r", encoding="utf-8") as json_file:
        payload = json.load(json_file)

    if isinstance(payload, dict):
        if "data" not in payload or not isinstance(payload["data"], list):
            raise DatasetValidationError("JSON object payload must include a data array.")
        payload = payload["data"]

    if not isinstance(payload, list):
        raise DatasetValidationError("JSON file must contain an array of objects.")

    rows = []
    for item in payload:
        if isinstance(item, dict):
            rows.append(_normalize_raw_row(item))

    if not rows:
        raise DatasetValidationError("JSON data is empty or malformed.")
    return rows


def _parse_excel(file_path):
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise DatasetValidationError("Missing openpyxl dependency; cannot parse Excel.") from error

    workbook = load_workbook(filename=file_path, data_only=True, read_only=True)
    worksheet = workbook.active
    row_iter = worksheet.iter_rows(values_only=True)
    header_row = next(row_iter, None)
    if not header_row:
        return []

    headers = [str(item).strip() if item is not None else "" for item in header_row]
    rows = []
    for values in row_iter:
        if values is None:
            continue
        item = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            item[header] = values[idx] if idx < len(values) else None
        if item:
            rows.append(_normalize_raw_row(item))

    return rows


def _normalize_raw_row(row):
    normalized = {}
    for key, value in row.items():
        if key is None:
            continue
        field = str(key).strip()
        if not field:
            continue
        if isinstance(value, str):
            normalized[field] = value.strip()
        else:
            normalized[field] = value
    return normalized


def clean_rows(raw_rows):
    if not raw_rows:
        raise DatasetValidationError("Imported data is empty.")

    standardized_rows = [_normalize_raw_row(row) for row in raw_rows if row]
    if not standardized_rows:
        raise DatasetValidationError("Imported data is empty.")

    dedup_rows = _drop_duplicate_rows(standardized_rows)
    numeric_fields = detect_numeric_fields(dedup_rows)

    for row in dedup_rows:
        for key, value in list(row.items()):
            if value is None or (isinstance(value, str) and value == ""):
                row[key] = 0 if key in numeric_fields else "Unknown"
            elif key in numeric_fields:
                row[key] = _to_float(value, 0)

    return dedup_rows, numeric_fields


def _drop_duplicate_rows(rows):
    unique = OrderedDict()
    for row in rows:
        signature = json.dumps(row, sort_keys=True, ensure_ascii=False, default=str)
        unique[signature] = row
    return list(unique.values())


def detect_numeric_fields(rows):
    if not rows:
        return []

    keys = set()
    for row in rows:
        keys.update(row.keys())

    numeric_fields = []
    for key in sorted(keys):
        values = [row.get(key) for row in rows if row.get(key) not in (None, "")]
        if not values:
            continue
        parse_count = 0
        for value in values:
            if _is_number(value):
                parse_count += 1
        if parse_count / len(values) >= 0.8:
            numeric_fields.append(key)
    return numeric_fields


def normalize_mapping_fields(cleaned_rows, numeric_fields, value_field=None, hierarchy_fields=None):
    all_fields = list(cleaned_rows[0].keys()) if cleaned_rows else []
    if value_field:
        value_field = value_field.strip()
    if hierarchy_fields:
        hierarchy_fields = [field.strip() for field in hierarchy_fields if field and field.strip()]

    if not value_field:
        value_field = numeric_fields[-1] if numeric_fields else None
    if value_field and value_field not in all_fields:
        raise DatasetValidationError("The specified value_field does not exist.")

    if not hierarchy_fields:
        hierarchy_fields = _auto_detect_hierarchy_fields(
            rows=cleaned_rows,
            all_fields=all_fields,
            numeric_fields=numeric_fields,
            value_field=value_field,
        )

    if not hierarchy_fields:
        raise DatasetValidationError("Could not detect hierarchy fields. Provide at least one dimension field.")

    for field in hierarchy_fields:
        if field not in all_fields:
            raise DatasetValidationError("hierarchy_fields contains an invalid field: {0}".format(field))

    return value_field, hierarchy_fields


def _auto_detect_hierarchy_fields(rows, all_fields, numeric_fields, value_field):
    row_count = len(rows)
    if row_count <= 0:
        return []

    numeric_set = set(numeric_fields or [])
    candidates = []
    for field in all_fields:
        if field == value_field or field in numeric_set:
            continue

        values = [str(row.get(field) or "").strip() for row in rows]
        non_empty = [value for value in values if value]
        if not non_empty:
            continue
        if _looks_like_identifier_field(field, non_empty):
            continue

        unique_ratio = len(set(non_empty)) / len(non_empty)
        candidates.append(
            {
                "field": field,
                "unique_ratio": unique_ratio,
            }
        )

    if not candidates:
        return []

    candidates.sort(key=lambda item: item["unique_ratio"])
    selected = [item["field"] for item in candidates[:3]]
    return selected


def _looks_like_identifier_field(field_name, values):
    normalized_name = (field_name or "").strip().lower()
    if normalized_name in {"id", "uuid", "key", "code", "serial_no"}:
        return True
    if any(token in normalized_name for token in ("id", "uuid", "url", "link", "address")):
        return True

    sample = values[:120]
    if not sample:
        return False
    url_like_count = 0
    for value in sample:
        text = value.lower()
        if text.startswith("http://") or text.startswith("https://") or text.startswith("//"):
            url_like_count += 1
    return url_like_count >= max(3, int(len(sample) * 0.6))


def build_tree_nodes(cleaned_rows, hierarchy_fields, value_field):
    node_map = {}

    for row in cleaned_rows:
        metric = _to_float(row.get(value_field), 0) if value_field else 1.0
        parent_key = None
        path_parts = []
        level = 0

        for field_index, field in enumerate(hierarchy_fields):
            raw_text = str(row.get(field) or "Unknown").strip() or "Unknown"
            node_names = _split_hierarchy_value(raw_text)
            for segment_index, raw_node_name in enumerate(node_names):
                node_name = _normalize_hierarchy_node_name(raw_node_name)
                if path_parts and path_parts[-1] == node_name:
                    continue

                path_parts.append(node_name)
                node_key = "/".join(path_parts)
                node = node_map.get(node_key)
                if node is None:
                    node = {
                        "node_key": node_key,
                        "node_name": node_name,
                        "level": level,
                        "parent_key": parent_key,
                        "value_num": 0.0,
                        "raw_payload": _build_node_raw_payload(
                            field=field,
                            field_index=field_index,
                            segment_index=segment_index,
                            parent_key=parent_key,
                            path_parts=path_parts,
                        ),
                        "children": set(),
                        "is_abnormal": 0,
                    }
                    node_map[node_key] = node

                node["value_num"] += metric
                if parent_key:
                    node_map[parent_key]["children"].add(node_key)

                parent_key = node_key
                level += 1

    _mark_abnormal_leaf_nodes(node_map)
    return list(node_map.values())


def _build_node_raw_payload(field, field_index, segment_index, parent_key, path_parts):
    path_snapshot = list(path_parts)
    return {
        "source_field": field,
        "source_field_index": field_index,
        "source_segment_index": segment_index,
        "path_key": "/".join(path_snapshot),
        "path_labels": path_snapshot,
        "depth": max(len(path_snapshot) - 1, 0),
        "parent_key": parent_key,
    }


def _split_hierarchy_value(value):
    text = _normalize_hierarchy_text(value)
    if not text:
        return ["Unknown"]

    separators = ["/", ">", "|", "::"]
    for separator in separators:
        if separator in text:
            parts = [_normalize_hierarchy_node_name(item) for item in text.split(separator)]
            parts = [item for item in parts if item]
            if 2 <= len(parts) <= 8:
                return parts
    return [_normalize_hierarchy_node_name(text)]


def _normalize_hierarchy_text(value):
    text = str(value or "").strip()
    if not text:
        return ""

    normalized = text.replace("\\", "/")
    replacements = {
        "\uFF0F": "/",
        "\uFF5C": "|",
        "\uFF1E": ">",
        "\u2192": ">",
        "\u21D2": ">",
        "\u21A6": ">",
        "\u27F6": ">",
        "->": ">",
        "=>": ">",
    }
    for source, target in replacements.items():
        normalized = normalized.replace(source, target)

    while "//" in normalized:
        normalized = normalized.replace("//", "/")
    while "||" in normalized:
        normalized = normalized.replace("||", "|")
    while ">>" in normalized:
        normalized = normalized.replace(">>", ">")
    return normalized.strip()


def _normalize_hierarchy_node_name(value):
    text = str(value or "").replace("\u3000", " ").strip()
    if not text:
        return "Unknown"
    return " ".join(text.split())

def _mark_abnormal_leaf_nodes(node_map):
    leaf_nodes = [node for node in node_map.values() if not node["children"]]
    if not leaf_nodes:
        return

    values = sorted([node["value_num"] for node in leaf_nodes])
    if len(values) < 4:
        return

    q1 = _quantile(values, 0.25)
    q3 = _quantile(values, 0.75)
    iqr = q3 - q1
    lower = q1 - 1.5 * iqr
    upper = q3 + 1.5 * iqr

    for node in leaf_nodes:
        if node["value_num"] < lower or node["value_num"] > upper:
            node["is_abnormal"] = 1


def _quantile(sorted_values, percent):
    if not sorted_values:
        return 0.0
    if len(sorted_values) == 1:
        return float(sorted_values[0])

    position = (len(sorted_values) - 1) * percent
    lower_idx = int(math.floor(position))
    upper_idx = int(math.ceil(position))
    lower_val = sorted_values[lower_idx]
    upper_val = sorted_values[upper_idx]
    if lower_idx == upper_idx:
        return float(lower_val)
    ratio = position - lower_idx
    return float(lower_val + (upper_val - lower_val) * ratio)


def persist_dataset_with_records(
    user_id,
    name,
    source_type,
    file_path,
    value_field,
    hierarchy_fields,
    cleaned_rows,
    nodes,
):
    dataset = Dataset(
        user_id=user_id,
        name=name,
        source_type=source_type,
        file_path=file_path,
        field_meta={
            "value_field": value_field,
            "hierarchy_fields": hierarchy_fields,
            "fields": list(cleaned_rows[0].keys()) if cleaned_rows else [],
        },
        total_rows=len(cleaned_rows),
        clean_status="pending",
    )

    db.session.add(dataset)
    db.session.flush()

    key_to_id = {}
    sorted_nodes = sorted(nodes, key=lambda item: (item["level"], item["node_key"]))
    for node in sorted_nodes:
        parent_id = key_to_id.get(node["parent_key"]) if node["parent_key"] else None
        raw_payload = dict(node["raw_payload"] or {})
        raw_payload["parent_id"] = int(parent_id) if parent_id is not None else None
        record = DataRecord(
            dataset_id=dataset.id,
            parent_id=parent_id,
            node_key=node["node_key"],
            node_name=node["node_name"],
            level=node["level"],
            value_num=Decimal(str(round(node["value_num"], 4))),
            raw_payload=raw_payload,
            is_abnormal=node["is_abnormal"],
        )
        db.session.add(record)
        db.session.flush()
        key_to_id[node["node_key"]] = record.id

    dataset.clean_status = "done"
    db.session.commit()
    return dataset, len(sorted_nodes)


def get_dataset_for_user(dataset_id, user_id):
    dataset = Dataset.query.filter_by(id=dataset_id, user_id=user_id).first()
    if dataset is None:
        raise DatasetValidationError("Dataset does not exist or is not accessible.")
    return dataset


def preview_dataset_rows(dataset, limit=20):
    file_path = resolve_dataset_file_path(dataset.file_path)
    if not file_path or not os.path.exists(file_path):
        raise DatasetValidationError("Dataset file does not exist and cannot be previewed.")

    rows = parse_rows(file_path=file_path, source_type=dataset.source_type)
    if not rows:
        return {
            "columns": [],
            "rows": [],
            "total_rows": 0,
            "preview_count": 0,
        }

    columns = []
    field_meta = dataset.field_meta if isinstance(dataset.field_meta, dict) else {}
    meta_fields = field_meta.get("fields") if isinstance(field_meta.get("fields"), list) else []
    for field in meta_fields:
        if field in rows[0]:
            columns.append(field)

    if not columns:
        columns = list(rows[0].keys())

    preview_rows = []
    for row in rows[:limit]:
        preview_row = {}
        for column in columns:
            value = row.get(column)
            preview_row[column] = "" if value is None else value
        preview_rows.append(preview_row)

    return {
        "columns": columns,
        "rows": preview_rows,
        "total_rows": len(rows),
        "preview_count": len(preview_rows),
    }


def preview_processed_records(dataset, limit=20):
    records = (
        DataRecord.query.filter_by(dataset_id=dataset.id)
        .order_by(DataRecord.id.asc())
        .all()
    )
    total_rows = len(records)

    record_map = {int(record.id): record for record in records}
    children_map = {}
    for record in records:
        parent_id = int(record.parent_id) if record.parent_id is not None else None
        children_map.setdefault(parent_id, []).append(record)

    for items in children_map.values():
        items.sort(
            key=lambda item: (
                int(item.level or 0),
                str(item.node_key or ""),
                int(item.id),
            )
        )

    ordered_records = []

    def visit(parent_id):
        for child in children_map.get(parent_id, []):
            ordered_records.append(child)
            visit(int(child.id))

    visit(None)
    records = ordered_records[:limit]

    columns = [
        "id",
        "parent_id",
        "parent_node_name",
        "node_key",
        "node_name",
        "level",
        "value_num",
        "parent_key",
        "path_key",
        "source_field",
    ]
    rows = []
    for record in records:
        raw_payload = record.raw_payload if isinstance(record.raw_payload, dict) else {}
        rows.append(
            {
                "id": int(record.id),
                "parent_id": int(record.parent_id) if record.parent_id is not None else None,
                "parent_node_name": (
                    record_map[int(record.parent_id)].node_name
                    if record.parent_id is not None and int(record.parent_id) in record_map
                    else ""
                ),
                "node_key": record.node_key,
                "node_name": record.node_name,
                "level": int(record.level or 0),
                "value_num": float(record.value_num or 0),
                "parent_key": raw_payload.get("parent_key") or "",
                "path_key": raw_payload.get("path_key") or record.node_key,
                "source_field": raw_payload.get("source_field") or "",
            }
        )

    return {
        "columns": columns,
        "rows": rows,
        "total_rows": total_rows,
        "preview_count": len(rows),
    }


def delete_dataset(dataset):
    cleanup_file_if_exists(dataset.file_path)
    db.session.delete(dataset)
    db.session.commit()


def resolve_dataset_file_path(file_path):
    if not file_path:
        return ""

    source_path = Path(file_path)
    if source_path.is_absolute():
        return str(source_path)

    candidates = [source_path]
    candidates.append(Path(current_app.root_path).parent / source_path)

    upload_root = Path(current_app.config.get("UPLOAD_ROOT", "uploads"))
    if not upload_root.is_absolute():
        upload_root = Path(current_app.root_path).parent / upload_root
    candidates.append(upload_root / source_path)

    for candidate in candidates:
        if candidate.exists():
            return str(candidate)

    return str(candidates[-1])


def cleanup_file_if_exists(file_path):
    try:
        resolved_path = resolve_dataset_file_path(file_path)
        if resolved_path and os.path.exists(resolved_path):
            os.remove(resolved_path)
    except OSError:
        pass


def _is_number(value):
    if isinstance(value, (int, float)):
        return True
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return False
        try:
            float(text)
            return True
        except ValueError:
            return False
    return False


def _to_float(value, default_value=0):
    if value in (None, ""):
        return float(default_value)
    try:
        return float(value)
    except (TypeError, ValueError):
        return float(default_value)
